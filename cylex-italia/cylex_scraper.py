"""
Scraper Cylex Italia con interfaccia grafica Tkinter.
Funzionalità:
  - Selezione geografica: Regione → Provincia → Città (o intera provincia)
  - Ricerca automatica di TUTTE le pagine per ogni città
  - Estrazione automatica dell'email dalla pagina di dettaglio azienda
  - Pulsante Stop per interrompere la ricerca
  - Pulsante Apri cartella per aprire Esplora File
  - Ordinamento alfabetico A→Z nel file Excel
  - Download automatico del database comuni in locale (una sola volta)

Uso:
    python cylex_scraper.py
"""
import csv
import os
import sys
import re
import time
import threading
import subprocess
import urllib.request
from datetime import datetime

import tkinter as tk
from tkinter import ttk, filedialog, messagebox

from curl_cffi import requests
from bs4 import BeautifulSoup
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment


# ============================================================
#  METADATI APPLICAZIONE
# ============================================================
APP_NAME = "Cylex Scraper"
APP_VERSION = "2.1"
APP_AUTHOR = "A. Santoro"
FOOTER_TEXT = f"{APP_NAME} {APP_VERSION} – by {APP_AUTHOR}"

# ============================================================
#  COSTANTI SCRAPING
# ============================================================
BASE_URL = "https://www.cylex-italia.it"
SEARCH_URL = BASE_URL + "/s"

HEADERS = {
    "Accept": "text/html,application/xhtml+xml,application/xml;q=0.9,image/avif,image/webp,*/*;q=0.8",
    "Accept-Language": "it-IT,it;q=0.9,en-US;q=0.8,en;q=0.7",
    "Accept-Encoding": "gzip, deflate, br",
    "Connection": "keep-alive",
    "Upgrade-Insecure-Requests": "1",
    "Sec-Fetch-Dest": "document",
    "Sec-Fetch-Mode": "navigate",
    "Sec-Fetch-Site": "none",
    "Sec-Fetch-User": "?1",
    "Cache-Control": "max-age=0",
}

RISULTATI_PER_PAGINA = 20

# ============================================================
#  DATABASE COMUNI ITALIANI
# ============================================================
CSV_URL = (
    "https://raw.githubusercontent.com/DarioCorno/database_comuni_italiani/"
    "main/csv/gi_comuni_cap.csv"
)
NOME_FILE_CSV = "gi_comuni_cap.csv"


def _cartella_base():
    if getattr(sys, "frozen", False):
        return os.path.dirname(sys.executable)
    return os.path.dirname(os.path.abspath(__file__))


def percorso_csv() -> str:
    return os.path.join(_cartella_base(), NOME_FILE_CSV)


def database_presente() -> bool:
    return os.path.exists(percorso_csv())


def scarica_database(progress_callback=None, forza=False):
    percorso = percorso_csv()

    if os.path.exists(percorso) and not forza:
        if progress_callback:
            progress_callback(100, "Database già presente in locale.")
        return percorso

    if progress_callback:
        progress_callback(0, "Download del database in corso...")

    def _reporthook(block_num, block_size, total_size):
        if total_size > 0 and progress_callback:
            percentuale = min(100, int(block_num * block_size * 100 / total_size))
            progress_callback(percentuale, f"Download: {percentuale}%")

    try:
        urllib.request.urlretrieve(CSV_URL, percorso, reporthook=_reporthook)
    except Exception as e:
        if os.path.exists(percorso):
            try:
                os.remove(percorso)
            except OSError:
                pass
        raise RuntimeError(f"Impossibile scaricare il database: {e}")

    if progress_callback:
        progress_callback(100, "Database scaricato con successo.")

    return percorso


def carica_comuni():
    percorso = percorso_csv()
    if not os.path.exists(percorso):
        raise FileNotFoundError(
            f"Database non trovato in {percorso}. Esegui prima scarica_database()."
        )

    comuni = []
    visti = set()
    with open(percorso, encoding="utf-8-sig") as f:
        reader = csv.DictReader(f, delimiter=";")
        for row in reader:
            codice = (row.get("codice_istat") or "").strip()
            if not codice or codice in visti:
                continue
            visti.add(codice)
            comuni.append({
                "codice_istat": codice,
                "comune": (row.get("denominazione_ita") or "").strip(),
                "provincia": (row.get("denominazione_provincia") or "").strip(),
                "sigla_provincia": (row.get("sigla_provincia") or "").strip(),
                "regione": (row.get("denominazione_regione") or "").strip(),
            })
    return comuni


def elenco_regioni(comuni):
    return sorted({c["regione"] for c in comuni if c["regione"]})


def province_per_regione(comuni, regione):
    return sorted({c["provincia"] for c in comuni
                   if c["regione"] == regione and c["provincia"]})


def comuni_per_provincia(comuni, provincia):
    return sorted({c["comune"] for c in comuni
                   if c["provincia"] == provincia and c["comune"]})


# ============================================================
#  FUNZIONI DI SCRAPING
# ============================================================

def cerca_pagina(categoria: str, citta: str, pagina: int = 1):
    params = {"q": categoria, "c": citta, "z": "", "p": pagina,
              "dst": "", "sUrl": "", "cUrl": ""}
    r = requests.get(SEARCH_URL, params=params, headers=HEADERS,
                     timeout=20, impersonate="chrome124")
    r.raise_for_status()
    return r.text


def rileva_totale_risultati(html: str):
    soup = BeautifulSoup(html, "html.parser")
    span = soup.select_one("div.lm-h h2 span.bold.text-muted")
    if span:
        match = re.search(r"su\s+(\d+)", span.get_text(strip=True))
        if match:
            return int(match.group(1))
    match = re.search(r"Risultati\s+\d+\s*-\s*\d+\s+su\s+(\d+)", soup.get_text())
    if match:
        return int(match.group(1))
    return None


def calcola_pagine(totali: int, per_pagina: int = RISULTATI_PER_PAGINA) -> int:
    if totali <= 0:
        return 1
    return (totali + per_pagina - 1) // per_pagina


def estrai_risultati(html: str):
    soup = BeautifulSoup(html, "html.parser")
    risultati = []
    for card in soup.select("div.lm-comp"):
        nome_tag = card.select_one("div.h4 a")
        nome = nome_tag.get_text(strip=True) if nome_tag else None
        indirizzo_tag = card.select_one("div.addr")
        indirizzo = indirizzo_tag.get_text(" ", strip=True) if indirizzo_tag else None
        telefono_tag = card.select_one("div.lm-ph span")
        telefono = telefono_tag.get_text(strip=True) if telefono_tag else None
        link_tag = card.select_one("a[href]")
        link = link_tag["href"] if link_tag else None
        if link and not link.startswith("http"):
            link = BASE_URL + link
        if nome:
            risultati.append({
                "nome": nome, "indirizzo": indirizzo,
                "telefono": telefono, "link": link
            })
    return risultati


def _decodifica_cf_email(hex_string: str):
    """
    Decodifica un indirizzo email protetto da Cloudflare Email Protection.
    Il formato è: cdn-cgi/l/email-protection#<hex>, dove il primo byte hex
    è la chiave XOR e i successivi sono i byte dell'email cifrati.
    """
    try:
        data = bytes.fromhex(hex_string)
        key = data[0]
        decoded = bytes(b ^ key for b in data[1:])
        return decoded.decode("utf-8")
    except Exception:
        return None


def _trova_email_in_html(soup: BeautifulSoup, html: str):
    """
    Cerca un'email nella pagina provando, in ordine:
      1) link mailto: dentro #cp-email
      2) qualsiasi link mailto: nella pagina
      3) link/span di Cloudflare Email Protection (#cp-email o generico)
      4) pattern email libero nell'HTML grezzo (fallback finale)
    """
    # 1) mailto specifico
    mailto_tag = soup.select_one("#cp-email a[href^='mailto:']")
    if not mailto_tag:
        # 2) mailto generico
        mailto_tag = soup.select_one("a[href^='mailto:']")
    if mailto_tag and mailto_tag.get("href"):
        email = mailto_tag["href"].replace("mailto:", "").split("?")[0].strip()
        if email:
            return email

    # 3) Cloudflare email protection: <a href="/cdn-cgi/l/email-protection#hex">
    #    oppure <span class="__cf_email__" data-cfemail="hex">
    cf_tag = soup.select_one("#cp-email a[href*='cdn-cgi/l/email-protection']")
    if not cf_tag:
        cf_tag = soup.select_one("a[href*='cdn-cgi/l/email-protection']")
    if cf_tag and cf_tag.get("href"):
        hex_part = cf_tag["href"].split("#", 1)[-1]
        email = _decodifica_cf_email(hex_part)
        if email:
            return email

    cf_span = soup.select_one("[data-cfemail]")
    if cf_span and cf_span.get("data-cfemail"):
        email = _decodifica_cf_email(cf_span["data-cfemail"])
        if email:
            return email

    # Anche cercando direttamente nel testo grezzo, nel caso il tag non sia
    # stato individuato dal parser (es. dentro markdown/attributi particolari)
    match_cf = re.search(r"email-protection#([0-9a-fA-F]+)", html)
    if match_cf:
        email = _decodifica_cf_email(match_cf.group(1))
        if email:
            return email

    match_cf2 = re.search(r'data-cfemail=["\']([0-9a-fA-F]+)["\']', html)
    if match_cf2:
        email = _decodifica_cf_email(match_cf2.group(1))
        if email:
            return email

    # 4) fallback finale: pattern email in chiaro nell'HTML
    match = RE_EMAIL.search(html)
    if match:
        return match.group(0)

    return None


def estrai_email_da_dettaglio(link: str, session=None, log_callback=None):
    """
    Visita la pagina di dettaglio dell'azienda ed estrae l'indirizzo email,
    gestendo anche le email offuscate con Cloudflare Email Protection.
    Restituisce None se non trovata o in caso di errore.
    """
    if not link:
        return None

    client = session or requests

    try:
        r = client.get(link, headers=HEADERS, timeout=25, impersonate="chrome124")
        r.raise_for_status()
    except Exception as e:
        if log_callback:
            log_callback(f"       ⚠️ Errore apertura pagina dettaglio: {e}")
        return None

    html = r.text
    soup = BeautifulSoup(html, "html.parser")

    email = _trova_email_in_html(soup, html)
    if email:
        return email

    if log_callback:
        log_callback("       ⚠️ Nessuna email trovata nella pagina di dettaglio.")
    return None


def salva_excel(nome_file: str, righe: list, categoria: str,
                etichetta_geografica: str, ricerche_effettuate: int,
                pagine_totali: int, multicitta: bool):
    """
    Salva il file Excel. Se multicitta=True, aggiunge le colonne
    'Città' e 'Provincia' e ordina per Città poi Nome.
    Altrimenti ordina solo per Nome (A→Z).
    Include sempre le colonne Email e Link.
    """
    wb = Workbook()
    ws = wb.active
    ws.title = "Risultati"

    if multicitta:
        n_colonne = 7
        intestazioni = ["Città", "Provincia", "Nome", "Indirizzo", "Telefono", "Email", "Link"]
    else:
        n_colonne = 5
        intestazioni = ["Nome", "Indirizzo", "Telefono", "Email", "Link"]

    # --- Riga 1: titolo ---
    titolo = f"Ricerca: {categoria} in {etichetta_geografica}"
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=n_colonne)
    cella_titolo = ws.cell(row=1, column=1, value=titolo)
    cella_titolo.font = Font(bold=True, size=16, color="FFFFFF")
    cella_titolo.fill = PatternFill(start_color="1F4E5F", end_color="1F4E5F", fill_type="solid")
    cella_titolo.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 30

    # --- Riga 2: sottotitolo ---
    data_ita = datetime.now().strftime("%d/%m/%Y %H:%M:%S")
    sottotitolo = (f"Data estrazione: {data_ita}  |  Ricerche effettuate: {ricerche_effettuate}  |  "
                   f"Pagine totali: {pagine_totali}")
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=n_colonne)
    cella_sub = ws.cell(row=2, column=1, value=sottotitolo)
    cella_sub.font = Font(italic=True, size=11, color="1F4E5F")
    cella_sub.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[2].height = 22

    # --- Riga 3: intestazioni ---
    ws.append(intestazioni)
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="2D7681", end_color="2D7681", fill_type="solid")
    header_align = Alignment(horizontal="center", vertical="center")
    for col_num in range(1, n_colonne + 1):
        cell = ws.cell(row=3, column=col_num)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_align

    # --- Dati ordinati ---
    if multicitta:
        righe_ordinate = sorted(righe, key=lambda r: (
            (r.get("citta") or "").lower(),
            (r.get("nome") or "").lower()
        ))
        for r in righe_ordinate:
            ws.append([r["citta"], r["provincia"], r["nome"],
                       r["indirizzo"], r["telefono"], r.get("email"), r["link"]])
        larghezze = {"A": 25, "B": 30, "C": 45, "D": 55, "E": 20, "F": 35, "G": 60}
    else:
        righe_ordinate = sorted(righe, key=lambda r: (r["nome"] or "").lower())
        for r in righe_ordinate:
            ws.append([r["nome"], r["indirizzo"], r["telefono"], r.get("email"), r["link"]])
        larghezze = {"A": 45, "B": 55, "C": 20, "D": 35, "E": 60}

    for col, larg in larghezze.items():
        ws.column_dimensions[col].width = larg

    ws.freeze_panes = "A4"
    ultima_colonna = chr(64 + n_colonne)
    ws.auto_filter.ref = f"A3:{ultima_colonna}{ws.max_row}"

    wb.save(nome_file)


# ============================================================
#  FUNZIONE DI RICERCA (thread separato)
# ============================================================

def esegui_ricerca(categoria, lista_citta, cartella_output, stop_event,
                   log_callback, fine_callback, etichetta_geografica,
                   estrai_email=True):
    """
    Esegue la ricerca per una o più città.
    lista_citta: lista di tuple (citta, provincia).
    Se estrai_email=True, visita la pagina di dettaglio di ogni azienda
    trovata per recuperarne l'indirizzo email.
    """
    try:
        def pulisci(s):
            s = s.lower().strip()
            s = re.sub(r"[^\w\-]+", "_", s)
            s = re.sub(r"_+", "_", s).strip("_")
            return s or "ricerca"

        cat = pulisci(categoria)
        etichetta_file = pulisci(etichetta_geografica)
        timestamp = datetime.now().strftime("%d-%m-%Y_%H-%M-%S")

        multicitta = len(lista_citta) > 1
        if multicitta:
            nome_file = f"cylex_{cat}_{etichetta_file}_{timestamp}.xlsx"
        else:
            nome_file = f"cylex_{cat}_{pulisci(lista_citta[0][0])}_{timestamp}.xlsx"

        percorso = os.path.join(cartella_output, nome_file)
        log_callback(f"[i] File di output: {percorso}")
        log_callback(f"[i] Città da elaborare: {len(lista_citta)}")

        tutti = []
        pagine_totali = 0
        ricerche_effettuate = 0
        interrotto = False

        # Sessione persistente (riusa cookie/TLS fingerprint) per le pagine di dettaglio
        session = requests.Session() if estrai_email else None

        for idx, (citta, provincia) in enumerate(lista_citta, 1):
            if stop_event.is_set():
                log_callback("⏹️  Interruzione richiesta dall'utente.")
                interrotto = True
                break

            log_callback(f"\n[{idx}/{len(lista_citta)}] Ricerca per: {citta} ({provincia})")

            try:
                html_prima = cerca_pagina(categoria, citta, 1)
            except Exception as e:
                log_callback(f"    ⚠️ Errore per {citta}: {e} - salto.")
                continue

            totale = rileva_totale_risultati(html_prima)
            if totale is None:
                log_callback(f"    ⚠️ Totale non rilevato. Salto.")
                continue

            pagine_da_scaricare = calcola_pagine(totale)
            log_callback(f"    Totale risultati: {totale} → {pagine_da_scaricare} pagine.")

            for p in range(1, pagine_da_scaricare + 1):
                if stop_event.is_set():
                    log_callback("⏹️  Interruzione richiesta dall'utente.")
                    interrotto = True
                    break

                if p == 1:
                    html = html_prima
                else:
                    try:
                        html = cerca_pagina(categoria, citta, p)
                    except Exception as e:
                        log_callback(f"    ⚠️ Errore pagina {p}: {e}")
                        break

                risultati = estrai_risultati(html)
                if not risultati:
                    break

                for r in risultati:
                    r["citta"] = citta
                    r["provincia"] = provincia
                    r["email"] = None

                    if estrai_email:
                        if stop_event.is_set():
                            interrotto = True
                            break
                        try:
                            r["email"] = estrai_email_da_dettaglio(
                                r["link"], session=session, log_callback=log_callback
                            )
                        except Exception as e:
                            log_callback(f"       ⚠️ Errore email per {r['nome']}: {e}")
                            r["email"] = None
                        if r["email"]:
                            log_callback(f"       ✉️  {r['nome']}: {r['email']}")
                        else:
                            log_callback(f"       ✉️  {r['nome']}: email non trovata")
                        # piccola pausa per non martellare il sito
                        time.sleep(0.5)

                if interrotto:
                    break

                tutti.extend(risultati)
                pagine_totali += 1

                # Pausa interrompibile tra pagine (3 s)
                if p < pagine_da_scaricare:
                    for _ in range(15):
                        if stop_event.is_set():
                            interrotto = True
                            break
                        time.sleep(0.2)
                    if interrotto:
                        break

            if interrotto:
                break

            ricerche_effettuate += 1
            log_callback(f"    ✓ {citta}: completata.")

            # Pausa tra città (3 s)
            if idx < len(lista_citta) and not interrotto:
                for _ in range(15):
                    if stop_event.is_set():
                        interrotto = True
                        break
                    time.sleep(0.2)

        # --- Salvataggio ---
        if tutti:
            salva_excel(
                percorso, tutti, categoria, etichetta_geografica,
                ricerche_effettuate, pagine_totali, multicitta
            )
            if interrotto:
                msg = (f"Ricerca interrotta.\n"
                       f"Salvati {len(tutti)} risultati ({ricerche_effettuate} città) in:\n{percorso}")
                log_callback(f"\n⏹️  {msg}")
                fine_callback(True, msg)
            else:
                msg = f"Salvati {len(tutti)} risultati ({ricerche_effettuate} città) in:\n{percorso}"
                log_callback(f"\n[✓] {msg}")
                fine_callback(True, msg)
        else:
            log_callback("[!] Nessun risultato trovato.")
            fine_callback(False, "Nessun risultato trovato.")

    except Exception as e:
        log_callback(f"❌ Errore: {e}")
        fine_callback(False, f"Errore: {e}")


# ============================================================
#  INTERFACCIA GRAFICA
# ============================================================

class CylexScraperGUI:
    def __init__(self, root):
        self.root = root
        self.root.title(f"{APP_NAME} {APP_VERSION}")
        self.root.geometry("760x760")
        self.root.minsize(720, 680)

        # Attributi DB
        self.comuni = []
        self.regioni = []

        # Variabili
        self.var_categoria = tk.StringVar()
        self.var_regione = tk.StringVar()
        self.var_provincia = tk.StringVar()
        self.var_citta = tk.StringVar()
        self.var_cartella = tk.StringVar(value=os.path.abspath("."))
        self.var_estrai_email = tk.BooleanVar(value=True)

        self.stop_event = threading.Event()
        self.thread_attivo = False

        self._costruisci_interfaccia()

        # Preparazione database dopo la costruzione della GUI
        self.root.after(100, self._prepara_database)

    def _costruisci_interfaccia(self):
        style = ttk.Style()
        style.theme_use("clam")
        style.configure("TLabel", font=("Segoe UI", 10))
        style.configure("TButton", font=("Segoe UI", 10))
        style.configure("Header.TLabel", font=("Segoe UI", 15, "bold"), foreground="#1F4E5F")

        main = ttk.Frame(self.root, padding=15)
        main.pack(fill="both", expand=True)

        # Titolo
        ttk.Label(main, text=f"🔍 {APP_NAME} {APP_VERSION}", style="Header.TLabel").grid(
            row=0, column=0, columnspan=3, sticky="w", pady=(0, 15))

        # Categoria
        ttk.Label(main, text="Categoria / Parole chiave:").grid(
            row=1, column=0, sticky="w", pady=5)
        ttk.Entry(main, textvariable=self.var_categoria, width=45).grid(
            row=1, column=1, columnspan=2, sticky="ew", pady=5)

        # Regione
        ttk.Label(main, text="Regione:").grid(row=2, column=0, sticky="w", pady=5)
        self.combo_regione = ttk.Combobox(main, textvariable=self.var_regione,
                                          state="readonly", width=43)
        self.combo_regione.grid(row=2, column=1, columnspan=2, sticky="ew", pady=5)
        self.combo_regione.bind("<<ComboboxSelected>>", self._on_regione_changed)

        # Provincia
        ttk.Label(main, text="Provincia:").grid(row=3, column=0, sticky="w", pady=5)
        self.combo_provincia = ttk.Combobox(main, textvariable=self.var_provincia,
                                            state="readonly", width=43)
        self.combo_provincia.grid(row=3, column=1, columnspan=2, sticky="ew", pady=5)
        self.combo_provincia.bind("<<ComboboxSelected>>", self._on_provincia_changed)

        # Città
        ttk.Label(main, text="Città:").grid(row=4, column=0, sticky="w", pady=5)
        self.combo_citta = ttk.Combobox(main, textvariable=self.var_citta,
                                        state="readonly", width=43)
        self.combo_citta.grid(row=4, column=1, columnspan=2, sticky="ew", pady=5)

        # Info
        ttk.Label(
            main,
            text="ℹ️  Seleziona 'Tutte le città della provincia' per elaborare l'intera provincia.",
            foreground="#1F4E5F", font=("Segoe UI", 9, "italic")
        ).grid(row=5, column=0, columnspan=3, sticky="w", pady=(5, 5))

        # Estrazione email
        ttk.Checkbutton(
            main, text="Estrai anche l'email dalla pagina di dettaglio (più lento)",
            variable=self.var_estrai_email
        ).grid(row=6, column=0, columnspan=3, sticky="w", pady=(0, 10))

        # Cartella output
        ttk.Label(main, text="Cartella di salvataggio:").grid(
            row=7, column=0, sticky="w", pady=5)
        frame_cart = ttk.Frame(main)
        frame_cart.grid(row=7, column=1, columnspan=2, sticky="ew", pady=5)
        ttk.Entry(frame_cart, textvariable=self.var_cartella).pack(
            side="left", fill="x", expand=True)
        ttk.Button(frame_cart, text="Sfoglia...", command=self._scegli_cartella).pack(
            side="left", padx=(5, 0))

        # Pulsanti azione
        frame_btn = ttk.Frame(main)
        frame_btn.grid(row=8, column=0, columnspan=3, sticky="ew", pady=10)

        self.btn_avvia = ttk.Button(frame_btn, text="▶  Avvia ricerca",
                                    command=self._avvia)
        self.btn_avvia.pack(side="left", fill="x", expand=True, padx=(0, 5))

        self.btn_stop = ttk.Button(frame_btn, text="⏹  Stop", command=self._stop,
                                   state="disabled")
        self.btn_stop.pack(side="left", fill="x", expand=True, padx=5)

        self.btn_apri = ttk.Button(frame_btn, text="📁  Apri cartella",
                                   command=self._apri_cartella)
        self.btn_apri.pack(side="left", fill="x", expand=True, padx=(5, 0))

        # Log
        ttk.Label(main, text="Log:").grid(row=9, column=0, sticky="w")
        frame_log = ttk.Frame(main)
        frame_log.grid(row=10, column=0, columnspan=3, sticky="nsew", pady=5)

        self.txt_log = tk.Text(frame_log, height=16, wrap="word",
                               font=("Consolas", 9), bg="#F5F5F5")
        scrollbar = ttk.Scrollbar(frame_log, orient="vertical",
                                  command=self.txt_log.yview)
        self.txt_log.configure(yscrollcommand=scrollbar.set)
        self.txt_log.pack(side="left", fill="both", expand=True)
        scrollbar.pack(side="right", fill="y")
        self.txt_log.configure(state="disabled")

        # Barra inferiore
        frame_bottom = ttk.Frame(main)
        frame_bottom.grid(row=11, column=0, columnspan=3, sticky="ew", pady=5)
        ttk.Button(frame_bottom, text="Pulisci log",
                   command=self._pulisci_log).pack(side="left")
        self.status = ttk.Label(frame_bottom, text="Pronto.", foreground="#555")
        self.status.pack(side="right")

        # Footer
        footer = ttk.Frame(main)
        footer.grid(row=12, column=0, columnspan=3, sticky="ew", pady=(10, 0))
        ttk.Separator(footer, orient="horizontal").pack(fill="x", pady=(0, 5))
        ttk.Label(footer, text=FOOTER_TEXT, foreground="#888",
                  font=("Segoe UI", 8, "italic"), anchor="center").pack(fill="x")

        main.columnconfigure(1, weight=1)
        main.rowconfigure(10, weight=1)

    # ---------- Preparazione database ----------

    def _prepara_database(self):
        if database_presente():
            try:
                self.comuni = carica_comuni()
                self.regioni = elenco_regioni(self.comuni)
                self.combo_regione["values"] = self.regioni
                self._log(f"[i] Database comuni caricato: {len(self.comuni)} comuni.")
                return
            except Exception as e:
                messagebox.showwarning(
                    "Database non valido",
                    f"Impossibile leggere il database locale:\n{e}\n\n"
                    "Verrà riscaricato."
                )

        self._finestra_download_database()

    def _finestra_download_database(self):
        win = tk.Toplevel(self.root)
        win.title("Download database comuni")
        win.geometry("500x210")
        win.transient(self.root)
        win.grab_set()
        win.resizable(False, False)

        win.update_idletasks()
        x = self.root.winfo_x() + (self.root.winfo_width() - 500) // 2
        y = self.root.winfo_y() + (self.root.winfo_height() - 210) // 2
        win.geometry(f"+{x}+{y}")

        ttk.Label(win, text="Download del database dei comuni italiani",
                  font=("Segoe UI", 12, "bold"), foreground="#1F4E5F").pack(pady=(15, 5))

        lbl_stato = ttk.Label(win, text="Preparazione...", font=("Segoe UI", 10))
        lbl_stato.pack(pady=5)

        progress = ttk.Progressbar(win, length=440, mode="determinate", maximum=100)
        progress.pack(pady=10)

        lbl_perc = ttk.Label(win, text="0%", font=("Segoe UI", 9))
        lbl_perc.pack()

        risultato = {"errore": None}

        def aggiorna(percentuale, messaggio):
            def _update():
                progress["value"] = percentuale
                lbl_stato.configure(text=messaggio.split("\n")[0])
                lbl_perc.configure(text=f"{percentuale}%")
            self.root.after(0, _update)

        def esegui():
            try:
                scarica_database(progress_callback=aggiorna, forza=True)
                self.comuni = carica_comuni()
                self.regioni = elenco_regioni(self.comuni)
            except Exception as e:
                risultato["errore"] = str(e)
            finally:
                self.root.after(0, chiudi)

        def chiudi():
            try:
                win.grab_release()
            except Exception:
                pass
            win.destroy()
            if risultato["errore"]:
                messagebox.showerror(
                    "Errore download",
                    f"Impossibile scaricare il database dei comuni:\n"
                    f"{risultato['errore']}\n\n"
                    "Verifica la connessione internet e riavvia l'applicazione."
                )
                self.root.destroy()
            else:
                self.combo_regione["values"] = self.regioni
                self._log(f"[i] Database comuni scaricato e caricato: {len(self.comuni)} comuni.")

        threading.Thread(target=esegui, daemon=True).start()

    # ---------- Callbacks menu a tendina ----------

    def _on_regione_changed(self, event=None):
        regione = self.var_regione.get()
        province = province_per_regione(self.comuni, regione)
        self.combo_provincia["values"] = province
        self.var_provincia.set("")
        self.var_citta.set("")
        self.combo_citta["values"] = []

    def _on_provincia_changed(self, event=None):
        provincia = self.var_provincia.get()
        comuni = comuni_per_provincia(self.comuni, provincia)
        valori = ["Tutte le città della provincia"] + comuni
        self.combo_citta["values"] = valori
        self.var_citta.set("")

    # ---------- Azioni ----------

    def _scegli_cartella(self):
        cartella = filedialog.askdirectory(initialdir=self.var_cartella.get())
        if cartella:
            self.var_cartella.set(cartella)

    def _log(self, msg):
        def _append():
            self.txt_log.configure(state="normal")
            self.txt_log.insert("end", msg + "\n")
            self.txt_log.see("end")
            self.txt_log.configure(state="disabled")
        self.root.after(0, _append)

    def _pulisci_log(self):
        self.txt_log.configure(state="normal")
        self.txt_log.delete("1.0", "end")
        self.txt_log.configure(state="disabled")

    def _apri_cartella(self):
        cartella = self.var_cartella.get().strip()
        if not os.path.isdir(cartella):
            messagebox.showwarning("Attenzione", "La cartella non esiste.")
            return
        try:
            if sys.platform == "win32":
                os.startfile(cartella)
            elif sys.platform == "darwin":
                subprocess.Popen(["open", cartella])
            else:
                subprocess.Popen(["xdg-open", cartella])
            self._log(f"[i] Aperta cartella: {cartella}")
        except Exception as e:
            messagebox.showerror("Errore", f"Impossibile aprire la cartella:\n{e}")

    def _stop(self):
        if self.thread_attivo:
            self.stop_event.set()
            self.btn_stop.configure(state="disabled")
            self._log("⏹️  Richiesta di interruzione inviata...")
            self.status.configure(text="Interruzione in corso...", foreground="orange")

    def _avvia(self):
        categoria = self.var_categoria.get().strip()
        regione = self.var_regione.get().strip()
        provincia = self.var_provincia.get().strip()
        citta_sel = self.var_citta.get().strip()
        cartella = self.var_cartella.get().strip()
        estrai_email = self.var_estrai_email.get()

        if not self.comuni:
            messagebox.showwarning("Attenzione", "Il database dei comuni non è ancora pronto.")
            return
        if not categoria:
            messagebox.showwarning("Attenzione", "Inserisci una categoria o parole chiave.")
            return
        if not regione:
            messagebox.showwarning("Attenzione", "Seleziona una regione.")
            return
        if not provincia:
            messagebox.showwarning("Attenzione", "Seleziona una provincia.")
            return
        if not citta_sel:
            messagebox.showwarning("Attenzione",
                                   "Seleziona una città o 'Tutte le città della provincia'.")
            return
        if not os.path.isdir(cartella):
            messagebox.showwarning("Attenzione", "La cartella di salvataggio non esiste.")
            return

        if citta_sel == "Tutte le città della provincia":
            lista_citta = [(c, provincia) for c in comuni_per_provincia(self.comuni, provincia)]
            etichetta = provincia
        else:
            lista_citta = [(citta_sel, provincia)]
            etichetta = citta_sel

        if not lista_citta:
            messagebox.showwarning("Attenzione", "Nessuna città trovata per questa provincia.")
            return

        if len(lista_citta) > 50:
            risposta = messagebox.askyesno(
                "Conferma",
                f"Stai per elaborare {len(lista_citta)} città.\n"
                f"Potrebbero volerci diversi minuti"
                f"{' (anche di più, con estrazione email attiva)' if estrai_email else ''}.\n\n"
                f"Continuare?"
            )
            if not risposta:
                return

        self.stop_event.clear()
        self.thread_attivo = True
        self.btn_avvia.configure(state="disabled")
        self.btn_stop.configure(state="normal")
        self.status.configure(text="Ricerca in corso...", foreground="#1F4E5F")

        def fine(successo, messaggio):
            self.thread_attivo = False
            self.btn_avvia.configure(state="normal")
            self.btn_stop.configure(state="disabled")
            if successo:
                self.status.configure(text="Completato.", foreground="green")
                messagebox.showinfo("Completato", messaggio)
            else:
                self.status.configure(text="Errore.", foreground="red")
                messagebox.showerror("Errore", messaggio)

        thread = threading.Thread(
            target=esegui_ricerca,
            args=(categoria, lista_citta, cartella, self.stop_event,
                  self._log, fine, etichetta),
            kwargs={"estrai_email": estrai_email},
            daemon=True
        )
        thread.start()


# ============================================================
#  AVVIO
# ============================================================

if __name__ == "__main__":
    root = tk.Tk()
    app = CylexScraperGUI(root)
    root.mainloop()
